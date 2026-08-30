"""
Import parser for QuickBooks "Sales by Item Detail" style exports.

This format is structurally nothing like a flat spreadsheet: it's a nested
report with category/dimension group headers, subtotal rows, and only the
actual transaction rows (Type == "Invoice") carrying real data. Everything
this parser does is a judgment call about how to map that report onto our
Order model — each one is commented so it's easy to override if it guesses
wrong for a particular row.
"""
import re
from collections import Counter

import openpyxl
from django.db import transaction

from orders.excel_import import get_or_create_customer
from orders.models import BoxType, Customer, Order, PlyType

# Line items that are service/setup charges, not actual carton orders —
# skipped rather than imported as a "box" with zero dimensions.
NON_BOX_KEYWORDS = [
    "printing cost", "printing plate making", "die cost",
    "die cut making cost", "block making", "transport",
]

# Ordered box-type keyword -> category name. First match wins, so more
# specific terms are listed before generic ones.
BOX_TYPE_KEYWORDS = [
    ("die cut", "Die-cut"),
    ("ceiling panel", "Ceiling Panel"),
    ("cup setters", "Cup Setters"),
    ("dividers board", "Dividers Board"),
    ("board", "Board"),
    ("partition", "Partition"),
    ("inner carton", "Inner Carton"),
    ("rsc", "RSC"),
    ("glued carton", "Glued Carton"),
]
DEFAULT_BOX_TYPE = "Carton"

PLY_PATTERN = re.compile(r"(\d+)\s*ply", re.IGNORECASE)

# Matches a contiguous run of "number[unit] x number[unit] x number[unit]"
# — e.g. "01 Ft x 10Ft", "3 X 3 X 9 Inches", "1005x205x330 mm". Captures
# the whole span so each number's own unit can be read individually,
# since QuickBooks data sometimes tags a unit on every number ("Ft" on
# both sides of "01 Ft x 10Ft") rather than once at the end.
DIM_SPAN_PATTERN = re.compile(
    r"(?:\d+(?:\.\d+)?\s*(?:ft|feet|inches?|in|mm|cm)?\s*[xX]\s*)+"
    r"\d+(?:\.\d+)?\s*(?:ft|feet|inches?|in|mm|cm)?",
    re.IGNORECASE,
)
DIM_TOKEN_PATTERN = re.compile(
    r"(\d+(?:\.\d+)?)\s*(ft|feet|inches?|in|mm|cm)?", re.IGNORECASE
)

# Matches this factory's round/board notation: "150x(R) 35mm",
# "978x(R)552mm" — two numbers separated by "x(R)" instead of plain "x".
ROUND_PATTERN = re.compile(
    r"(\d+(?:\.\d+)?)\s*[xX]\s*\(R\)\s*(\d+(?:\.\d+)?)\s*(mm|cm|inches?|in|ft|feet)?",
    re.IGNORECASE,
)

UNIT_TO_CM = {
    "inch": 2.54, "inches": 2.54, "in": 2.54,
    "mm": 0.1, "cm": 1.0, "ft": 30.48, "feet": 30.48,
}


def _unit_factor(unit):
    key = (unit or "cm").lower().rstrip("es").rstrip("s")
    return UNIT_TO_CM.get(key) or UNIT_TO_CM.get((unit or "cm").lower()) or 1.0


def extract_dimensions_cm(*texts):
    """Tries each text in order, returns (length_cm, width_cm, height_cm)
    or None if no dimension pattern found in any of them. Missing third
    dimension (2D items like panels/partitions/round boards) gets
    height_cm = 0.

    Each number keeps its own unit when the source specifies one on every
    number (e.g. "01 Ft x 10Ft"); when only the last number carries a unit
    (e.g. "3 X 3 X 9 Inches"), that unit is applied back to the earlier
    numbers too, since that's the common convention in this data.
    """
    for text in texts:
        if not text:
            continue

        m = ROUND_PATTERN.search(text)
        if m:
            a, b, unit = m.groups()
            factor = _unit_factor(unit)
            return round(float(a) * factor, 2), round(float(b) * factor, 2), 0

        span_match = DIM_SPAN_PATTERN.search(text)
        if not span_match:
            continue
        tokens = DIM_TOKEN_PATTERN.findall(span_match.group(0))
        if len(tokens) < 2:
            continue

        # Backfill: if the last token has a unit and earlier ones don't,
        # apply the last unit to all of them (common in this data).
        last_unit = next((u for _, u in reversed(tokens) if u), None)
        values_cm = []
        for num, unit in tokens[:3]:
            factor = _unit_factor(unit or last_unit)
            values_cm.append(round(float(num) * factor, 2))

        while len(values_cm) < 3:
            values_cm.append(0)
        return tuple(values_cm)
    return None


def extract_box_type(memo):
    memo_lower = (memo or "").lower()
    for keyword, label in BOX_TYPE_KEYWORDS:
        if keyword in memo_lower:
            return label
    return DEFAULT_BOX_TYPE


def extract_ply(memo):
    m = PLY_PATTERN.search(memo or "")
    return f"{int(m.group(1))} Ply" if m else "Unspecified"


def extract_is_printed(memo):
    memo_lower = (memo or "").lower()
    if "unprinted" in memo_lower:
        return False
    if "printed" in memo_lower:
        return True
    return False  # no explicit signal — defaults to unprinted, flagged in the report


def is_non_box_line(memo, item_path):
    combined = f"{memo or ''} {item_path or ''}".lower()
    return any(kw in combined for kw in NON_BOX_KEYWORDS)


def _find_data_sheet(wb):
    """QuickBooks exports sometimes include an extra info/tips sheet before
    the real data sheet — don't assume the data is on the first sheet."""
    required = {"Type", "Num", "Memo", "Name", "Item", "Qty"}
    for ws in wb.worksheets:
        try:
            header_row = next(ws.iter_rows(min_row=1, max_row=1))
        except StopIteration:
            continue
        header_set = {str(c.value).strip() for c in header_row if c.value}
        if required.issubset(header_set):
            return ws
    return wb.worksheets[0]  # fallback: nothing matched, let the caller report missing columns


def looks_like_quickbooks_export(file_obj):
    """Quick signature check: does any sheet contain the QuickBooks column
    headers we expect, rather than our own flat format?"""
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
        ws = _find_data_sheet(wb)
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        file_obj.seek(0)
        header_set = {str(h).strip() for h in header if h}
        return {"Type", "Num", "Memo", "Name", "Item", "Qty"}.issubset(header_set)
    except Exception:
        try:
            file_obj.seek(0)
        except Exception:
            pass
        return False


def import_from_quickbooks_export(file_obj, dry_run=False):
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = _find_data_sheet(wb)
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: idx for idx, name in enumerate(header) if name}

    required = ["Type", "Date", "Num", "Memo", "Name", "Item", "Item Description", "Qty", "Sales Price"]
    missing = [c for c in required if c not in col]
    if missing:
        return {
            "created": 0, "skipped": 0, "bad_dates": 0, "skipped_rows": [],
            "missing_columns": missing, "non_box_skipped": 0, "no_dimension_skipped": 0,
            "no_printed_signal": 0, "split_invoice_lines": 0,
        }

    all_rows = list(ws.iter_rows(min_row=2, values_only=True))
    detail_rows = [r for r in all_rows if r[col["Type"]] == "Invoice"]

    # Count occurrences of each QuickBooks invoice number so we know which
    # ones need a "-1", "-2" line suffix (one QuickBooks invoice can cover
    # several different box types/products as separate line items — our
    # Order model is one-box-type-per-row, so each line needs its own
    # unique invoice_number while staying traceable to the original).
    num_counts = Counter(r[col["Num"]] for r in detail_rows if r[col["Num"]])
    line_seen = Counter()

    stats = {
        "created": 0, "skipped": 0, "bad_dates": 0, "skipped_rows": [],
        "missing_columns": [], "non_box_skipped": 0, "no_dimension_skipped": 0,
        "no_printed_signal": 0, "split_invoice_lines": sum(1 for v in num_counts.values() if v > 1),
    }

    try:
        with transaction.atomic():
            _import_detail_rows(detail_rows, col, num_counts, line_seen, stats)
            if dry_run:
                raise _DryRunRollback()
    except _DryRunRollback:
        pass

    return stats


class _DryRunRollback(Exception):
    pass


def _import_detail_rows(detail_rows, col, num_counts, line_seen, stats):
    for idx, row in enumerate(detail_rows, start=1):
        qb_num = row[col["Num"]]
        memo = row[col["Memo"]]
        item_path = row[col["Item"]]
        item_desc = row[col["Item Description"]]
        customer_name = row[col["Name"]]
        date_val = row[col["Date"]]
        qty = row[col["Qty"]]
        price = row[col["Sales Price"]]

        if not qb_num or not customer_name or not date_val:
            stats["skipped"] += 1
            stats["skipped_rows"].append((idx, "missing invoice number, customer, or date"))
            continue

        if is_non_box_line(memo, item_path):
            stats["non_box_skipped"] += 1
            stats["skipped_rows"].append(
                (idx, f"invoice {qb_num}: service/setup charge, not a box order — skipped")
            )
            continue

        dims = extract_dimensions_cm(item_desc, memo)
        if dims is None:
            stats["no_dimension_skipped"] += 1
            stats["skipped_rows"].append(
                (idx, f"invoice {qb_num}: couldn't find dimensions in '{item_desc or memo}' — skipped, add manually")
            )
            continue
        length_cm, width_cm, height_cm = dims

        # Build the invoice number our system will store: only append a
        # line suffix when this QuickBooks invoice number covers multiple
        # line items, so single-line invoices keep their original number.
        if num_counts[qb_num] > 1:
            line_seen[qb_num] += 1
            invoice_number = f"{qb_num}-{line_seen[qb_num]}"
        else:
            invoice_number = str(qb_num)

        memo_lower = (memo or "").lower()
        if "printed" not in memo_lower:  # covers neither "printed" nor "unprinted"
            stats["no_printed_signal"] += 1

        customer, _ = get_or_create_customer(str(customer_name).strip())
        box_type, _ = BoxType.objects.get_or_create(name=extract_box_type(memo))
        ply_type, _ = PlyType.objects.get_or_create(name=extract_ply(memo))

        order, was_created = Order.objects.get_or_create(
            invoice_number=invoice_number,
            defaults={
                "customer": customer,
                "box_type": box_type,
                "ply_type": ply_type,
                "length_cm": length_cm,
                "width_cm": width_cm,
                "height_cm": height_cm,
                "date_issued": date_val.date() if hasattr(date_val, "date") else date_val,
                "issued_price": price or 0,   # QuickBooks "Sales Price" = per-unit price
                "quantity": int(qty) if qty else 1,
                "is_printed": extract_is_printed(memo),
            },
        )
        if was_created:
            stats["created"] += 1
        else:
            stats["skipped"] += 1
            stats["skipped_rows"].append((idx, f"invoice {invoice_number} already exists"))
